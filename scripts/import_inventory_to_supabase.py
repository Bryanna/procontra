#!/usr/bin/env python3
"""Validate and idempotently import operational inventory into Supabase."""

from __future__ import annotations

import argparse
import json
import os
import urllib.parse
import urllib.request
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook


HEADERS = [
    "CODIGO", "PRODUCTO", "SUCURSAL", "EXISTENCIA", "INVENTARIO_MINIMO",
    "LOTE", "VENCIMIENTO", "COSTO", "PRECIO", "ACTUALIZADO_EN",
]
Transport = Callable[..., object]


def load_environment(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    if path.exists():
        for raw in path.read_text(encoding="utf-8").splitlines():
            if raw.strip() and not raw.lstrip().startswith("#") and "=" in raw:
                key, value = raw.split("=", 1)
                values[key.strip()] = value.strip()
    values.update({key: value for key, value in os.environ.items() if value})
    missing = [key for key in ("SUPABASE_URL", "SUPABASE_SERVICE_ROLE_KEY") if not values.get(key)]
    if missing:
        raise ValueError(f"Faltan variables requeridas: {', '.join(missing)}")
    return values


def iso_value(value: object) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def read_inventory_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if "Carga_Inventario" not in workbook.sheetnames:
        raise ValueError("El archivo requiere la hoja Carga_Inventario")
    sheet = workbook["Carga_Inventario"]
    values = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(values)]
    if headers != HEADERS:
        raise ValueError(f"Encabezados requeridos: {' | '.join(HEADERS)}")

    rows: list[dict] = []
    for row_number, raw_row in enumerate(values, start=2):
        if all(value is None or str(value).strip() == "" for value in raw_row):
            continue
        row = dict(zip(HEADERS, raw_row, strict=True))
        missing = [key for key, value in row.items() if value is None or str(value).strip() == ""]
        if missing:
            raise ValueError(f"Fila {row_number}: faltan {', '.join(missing)}")
        if not isinstance(row["CODIGO"], str):
            raise ValueError(f"Fila {row_number}: CODIGO debe ser texto para conservar ceros iniciales")
        row["CODIGO"] = row["CODIGO"].strip()
        row["SUCURSAL"] = str(row["SUCURSAL"]).strip().zfill(2)
        row["PRODUCTO"] = str(row["PRODUCTO"]).strip()
        row["LOTE"] = str(row["LOTE"]).strip()
        expiry = row["VENCIMIENTO"]
        if isinstance(expiry, datetime):
            expiry = expiry.date()
        row["VENCIMIENTO"] = iso_value(expiry)
        row["ACTUALIZADO_EN"] = iso_value(row["ACTUALIZADO_EN"])
        rows.append(row)
    return rows


def numeric(value: object, field: str) -> str:
    try:
        number = Decimal(str(value))
    except InvalidOperation as error:
        raise ValueError(f"{field} debe ser numérico") from error
    if number < 0:
        raise ValueError(f"{field} no puede ser negativo")
    return format(number.normalize(), "f")


def build_inventory_positions(
    rows: list[dict],
    product_ids: dict[str, str],
    branch_ids: dict[str, tuple[str, int]],
    source: str,
) -> list[dict]:
    positions: list[dict] = []
    seen: set[tuple[str, str, str]] = set()
    for row in rows:
        code = str(row["CODIGO"])
        branch = str(row["SUCURSAL"]).zfill(2)
        if code not in product_ids:
            raise ValueError(f"Producto desconocido: {code}")
        if branch not in branch_ids:
            raise ValueError(f"Sucursal desconocida: {branch}")
        key = (code, branch, str(row["LOTE"]))
        if key in seen:
            raise ValueError(f"Posición duplicada: {code}/{branch}/{row['LOTE']}")
        seen.add(key)
        branch_uuid, branch_tenant_id = branch_ids[branch]
        positions.append({
            "f_uuid_producto": product_ids[code],
            "f_uuid_sucursal": branch_uuid,
            "f_disponible": numeric(row["EXISTENCIA"], "EXISTENCIA"),
            "f_reservado": "0",
            "f_minimo_reorden": numeric(row["INVENTARIO_MINIMO"], "INVENTARIO_MINIMO"),
            "f_lote": str(row["LOTE"]),
            "f_fecha_vencimiento": str(row["VENCIMIENTO"]),
            "f_costo": numeric(row["COSTO"], "COSTO"),
            "f_precio": numeric(row["PRECIO"], "PRECIO"),
            "f_actualizado_en": str(row["ACTUALIZADO_EN"]),
            "f_fuente": source,
            "f_idsucursal": branch_tenant_id,
        })
    return positions


def request_json(url: str, service_key: str) -> list[dict]:
    request = urllib.request.Request(url, headers={
        "apikey": service_key,
        "Authorization": f"Bearer {service_key}",
        "Accept-Profile": "api",
    })
    with urllib.request.urlopen(request, timeout=120) as response:
        return json.loads(response.read())


def fetch_reference_map(supabase_url: str, service_key: str, table: str, key: str) -> dict[str, str]:
    output: dict[str, str] = {}
    offset = 0
    while True:
        query = urllib.parse.urlencode({"select": f"f_uuid,{key}", "limit": 1000, "offset": offset})
        rows = request_json(f"{supabase_url.rstrip('/')}/rest/v1/{table}?{query}", service_key)
        output.update({str(row[key]): str(row["f_uuid"]) for row in rows})
        if len(rows) < 1000:
            return output
        offset += 1000


def fetch_branch_map(supabase_url: str, service_key: str) -> dict[str, tuple[str, int]]:
    query = urllib.parse.urlencode({"select": "f_uuid,f_codigo,f_idsucursal", "limit": 1000})
    rows = request_json(f"{supabase_url.rstrip('/')}/rest/v1/t_sucursales?{query}", service_key)
    return {str(row["f_codigo"]): (str(row["f_uuid"]), int(row["f_idsucursal"])) for row in rows}


def upsert_positions(positions: list[dict], supabase_url: str, service_key: str, batch_size: int = 500) -> int:
    endpoint = f"{supabase_url.rstrip('/')}/rest/v1/t_posiciones_inventario?on_conflict=f_uuid_producto,f_uuid_sucursal,f_lote"
    batches = 0
    for start in range(0, len(positions), batch_size):
        request = urllib.request.Request(
            endpoint,
            data=json.dumps(positions[start:start + batch_size], ensure_ascii=False).encode("utf-8"),
            method="POST",
            headers={
                "Content-Type": "application/json",
                "apikey": service_key,
                "Authorization": f"Bearer {service_key}",
                "Prefer": "resolution=merge-duplicates,return=minimal",
                "Content-Profile": "api",
            },
        )
        with urllib.request.urlopen(request, timeout=120) as response:
            response.read()
        batches += 1
    return batches


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--env-file", type=Path, default=Path(".env.local"))
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    rows = read_inventory_rows(args.workbook)
    if args.dry_run or not rows:
        print(json.dumps({"rows": len(rows), "dryRun": True, "source": args.workbook.name}))
        return

    environment = load_environment(args.env_file)
    url = environment["SUPABASE_URL"]
    key = environment["SUPABASE_SERVICE_ROLE_KEY"]
    products = fetch_reference_map(url, key, "t_productos", "f_codigo")
    branches = fetch_branch_map(url, key)
    positions = build_inventory_positions(rows, products, branches, args.workbook.name)
    batch_count = upsert_positions(positions, url, key)
    print(json.dumps({"rows": len(positions), "batches": batch_count, "dryRun": False, "source": args.workbook.name}))


if __name__ == "__main__":
    main()
