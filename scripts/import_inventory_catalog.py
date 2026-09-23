#!/usr/bin/env python3
"""Import the pharmacy master workbook as a product catalog.

This importer deliberately does not create stock quantities. The source workbook
contains product codes/descriptions and an optional location, not branch stock.
"""

from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REQUIRED_HEADERS = ("CODIGO", "PRODUCTO")
OPERATIONAL_FIELDS = (
    "SUCURSAL",
    "EXISTENCIA",
    "INVENTARIO_MINIMO",
    "LOTE",
    "VENCIMIENTO",
    "COSTO",
    "PRECIO",
    "ACTUALIZADO_EN",
)


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def find_header_row(sheet: Any) -> tuple[int, list[str]]:
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        headers = [clean(value).upper() for value in row]
        if all(required in headers for required in REQUIRED_HEADERS):
            return row_number, headers
    raise ValueError("No se encontró una fila con los encabezados CODIGO y PRODUCTO")


def import_catalog(source: Path, output: Path) -> dict[str, Any]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    if "Datos" not in workbook.sheetnames:
        raise ValueError("El archivo debe contener la hoja Datos")

    sheet = workbook["Datos"]
    header_row, headers = find_header_row(sheet)
    positions = {header: index for index, header in enumerate(headers) if header}
    location_index = positions.get("UBICACION")

    products: list[dict[str, str]] = []
    seen_pairs: set[tuple[str, str]] = set()
    code_descriptions: dict[str, str] = {}
    exact_duplicates = 0
    code_conflicts = 0
    blank_codes = 0
    blank_descriptions = 0
    missing_location = 0
    source_rows = 0

    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(clean(value) for value in row):
            continue
        source_rows += 1
        code = clean(row[positions["CODIGO"]])
        description = clean(row[positions["PRODUCTO"]])
        location = clean(row[location_index]) if location_index is not None else ""

        if not code:
            blank_codes += 1
            continue
        if not description:
            blank_descriptions += 1
            continue

        pair = (code, description)
        if pair in seen_pairs:
            exact_duplicates += 1
            continue
        seen_pairs.add(pair)

        prior = code_descriptions.get(code)
        if prior is not None and prior != description:
            code_conflicts += 1
        else:
            code_descriptions[code] = description

        if not location:
            missing_location += 1
        products.append({"code": code, "name": description, "location": location})

    source_modified = datetime.fromtimestamp(source.stat().st_mtime, tz=timezone.utc).isoformat()
    available_headers = {header for header in headers if header}
    report: dict[str, Any] = {
        "sourceFile": source.name,
        "sourceSheet": "Datos",
        "sourceModifiedAt": source_modified,
        "headerRow": header_row,
        "sourceRows": source_rows,
        "catalogEntries": len(products),
        "uniqueCodes": len({product["code"] for product in products}),
        "exactDuplicatesIgnored": exact_duplicates,
        "codeDescriptionConflicts": code_conflicts,
        "blankCodesIgnored": blank_codes,
        "blankDescriptionsIgnored": blank_descriptions,
        "missingLocation": missing_location,
        "missingOperationalFields": [
            field for field in OPERATIONAL_FIELDS if field not in available_headers
        ],
        "stockStatus": "NOT_AVAILABLE",
    }
    payload = {"metadata": report, "products": products}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    print(json.dumps(import_catalog(args.source, args.output), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
